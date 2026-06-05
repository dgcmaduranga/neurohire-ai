import os
import uuid

import docx
import fitz
import pytesseract
from PIL import Image
from fastapi import HTTPException, UploadFile


UPLOAD_DIR = "app/uploads/resumes"
ALLOWED_EXTENSIONS = [".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg"]


def get_file_extension(filename: str) -> str:
    return os.path.splitext(filename.lower())[1]


def validate_file_type(filename: str):
    extension = get_file_extension(filename)

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Please upload PDF, DOCX, TXT, PNG, JPG, or JPEG.",
        )

    return extension


async def save_upload_file(
    file: UploadFile,
    upload_dir: str = UPLOAD_DIR,
) -> str:
    os.makedirs(upload_dir, exist_ok=True)

    original_filename = file.filename or "resume"
    extension = validate_file_type(original_filename)

    safe_name = original_filename.replace(" ", "_")
    unique_name = f"{uuid.uuid4().hex}_{safe_name}"

    file_path = os.path.join(upload_dir, unique_name)

    content = await file.read()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="Uploaded file is empty.",
        )

    with open(file_path, "wb") as f:
        f.write(content)

    return file_path


def extract_text_from_pdf(file_path: str) -> str:
    text_parts = []

    try:
        pdf = fitz.open(file_path)

        for page in pdf:
            page_text = page.get_text("text")

            if page_text:
                text_parts.append(page_text)

        if text_parts and len(" ".join(text_parts).strip()) >= 50:
            pdf.close()
            return "\n".join(text_parts).strip()

        ocr_text_parts = []

        for page_index in range(len(pdf)):
            page = pdf[page_index]
            pix = page.get_pixmap(dpi=200)

            image_path = f"{file_path}_page_{page_index}.png"
            pix.save(image_path)

            try:
                image_text = pytesseract.image_to_string(Image.open(image_path))
                if image_text.strip():
                    ocr_text_parts.append(image_text)
            finally:
                if os.path.exists(image_path):
                    os.remove(image_path)

        pdf.close()

        return "\n".join(ocr_text_parts).strip()

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"PDF text extraction failed: {str(e)}",
        )


def extract_text_from_docx(file_path: str) -> str:
    try:
        document = docx.Document(file_path)
        text_parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())

        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())

                if row_text:
                    text_parts.append(" | ".join(row_text))

        return "\n".join(text_parts).strip()

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"DOCX text extraction failed: {str(e)}",
        )


def extract_text_from_txt(file_path: str) -> str:
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
            return file.read().strip()

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"TXT text extraction failed: {str(e)}",
        )


def extract_text_from_image(file_path: str) -> str:
    try:
        image = Image.open(file_path)
        text = pytesseract.image_to_string(image)

        return text.strip()

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Image OCR failed: {str(e)}",
        )


def extract_resume_text(file_path: str, filename: str) -> str:
    extension = validate_file_type(filename)

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    if extension == ".txt":
        return extract_text_from_txt(file_path)

    if extension in [".png", ".jpg", ".jpeg"]:
        return extract_text_from_image(file_path)

    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Please upload PDF, DOCX, TXT, PNG, JPG, or JPEG.",
    )